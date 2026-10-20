"""Step 2: classify the step-1 loan request and fill the matching bank template with a maximum-loan estimate.
Step 3: turn the assessment into a draft reply to the customer.

Usage:
    python credit_assessment.py                 # newest <stamp>.md in OUTBOX_DIR
    python credit_assessment.py <file.md>       # a specific step-1 result
    python credit_assessment.py <outbox folder> # the step-1 result inside that folder
    ... --no-reply                              # assessment only
    ... --reply-only                            # regenerate the customer reply from an existing assessment

Every template file in TEMPLATES_DIR (default: ./templates, override with env CREDIT_TEMPLATES_DIR
or --templates) is sent with every request, whatever its name; the model classifies the request,
picks the matching template from its content and reports both in the first two lines of the answer.

Writes next to the source markdown:
    assessment_input.txt      - the exact step-2 prompt sent to the model
    <stamp>_assessment.md     - the filled template (internal, for the analyst)
    customer_reply_input.txt  - the exact step-3 prompt
    <stamp>_customer_reply.md - draft letter to the customer, for the advisor to complete and send
"""
from __future__ import annotations

import argparse
import logging
import os
import re
import sys
from datetime import datetime
from pathlib import Path

import read_inbox as pipeline
from assessment_prompt import LETTER_TYPES, LOAN_TYPES, build_assessment_prompt, build_customer_reply_prompt

ROOT = Path(__file__).resolve().parent
TEMPLATES_DIR = Path(os.getenv("CREDIT_TEMPLATES_DIR", ROOT / "templates"))
EXPECTED_TEMPLATES = 3  # one per loan type in assessment_prompt.LOAN_TYPES
ASSESSMENT_SUFFIX = "_assessment.md"
ASSESSMENT_MAX_TOKENS = 8192
REPLY_SUFFIX = "_customer_reply.md"
REPLY_MAX_TOKENS = 4096
LETTER_TYPE_LABELS = {
    "offer": "Offer",
    "reduced_offer": "Offer for a lower amount",
    "documents_request": "Request for documents",
    "needs_review": "Advisor decision needed (meeting invitation)",
}
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".xml", ".log"}
MAX_XLSX_ROWS = 2000
MARKER = re.compile(r"^(REQUEST_TYPE|TEMPLATE)\s*:\s*(.+?)$", re.IGNORECASE)

log = logging.getLogger("fomo.assessment")


# ---------------------------------------------------------------- Templates

def read_docx(path: Path) -> str:
    """Word -> markdown-like text that keeps headings, lists and tables, so the model can mirror the layout."""
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise RuntimeError("python-docx is not installed: pip install python-docx")

    def paragraph(p: Paragraph) -> str | None:
        text = p.text.strip()
        if not text:
            return None
        style = (p.style.name if p.style is not None else "").lower()
        level = re.search(r"(heading|otsikko|nagłówek)\s*(\d)", style)
        if level:
            return "#" * min(int(level.group(2)), 6) + " " + text
        if style == "title":
            return "# " + text
        if "list" in style or "luettelo" in style or "lista" in style or p._p.pPr is not None and p._p.pPr.numPr is not None:
            return "- " + text
        return text

    def table(t: Table) -> list[str]:
        rows = []
        for row in t.rows:
            cells, seen = [], set()
            for c in row.cells:
                # Merged cells are returned once per grid column; keep each physical cell once.
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                cells.append(" ".join(c.text.split()))
            rows.append("| " + " | ".join(cells) + " |")
        return rows

    doc = docx.Document(str(path))
    out = []
    # Walk the body in document order so tables stay next to the paragraphs that describe them.
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            line = paragraph(Paragraph(child, doc))
            if line:
                out.append(line)
        elif child.tag.endswith("}tbl"):
            out.extend(table(Table(child, doc)))
    return "\n".join(out)


def read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    values = load_workbook(path, data_only=True, read_only=True)
    formulas = load_workbook(path, data_only=False, read_only=True)
    out = []
    for ws_v, ws_f in zip(values.worksheets, formulas.worksheets):
        out.append(f"### Sheet: {ws_v.title}")
        for i, (row_v, row_f) in enumerate(zip(ws_v.iter_rows(), ws_f.iter_rows())):
            if i >= MAX_XLSX_ROWS:
                out.append(f"... (truncated after {MAX_XLSX_ROWS} rows)")
                break
            cells = []
            for cv, cf in zip(row_v, row_f):
                if cv.value is None and cf.value is None:
                    continue
                text = "" if cv.value is None else str(cv.value)
                # Formulas carry the bank's calculation rules, so keep them alongside the cached value.
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    text = f"{text} [{cf.value}]".strip()
                cells.append(f"{cv.coordinate}: {text}")
            if cells:
                out.append(" | ".join(cells))
    values.close()
    formulas.close()
    return "\n".join(out)


def read_template(path: Path) -> str | None:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace").strip()
    if suffix in {".html", ".htm"}:
        return pipeline.html_to_text(path.read_text(encoding="utf-8", errors="replace"))
    if suffix == ".docx":
        return read_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    if suffix in {".doc", ".xls"}:
        raise RuntimeError(f"{path.name}: old Office format, save it as .docx / .xlsx")
    # PDF (text layer + OCR of image pages) and images (OCR) via the step-1 readers.
    return pipeline.read_document(path)


def load_templates(folder: Path) -> list[tuple[str, str]]:
    folder = pipeline.long_path(folder)
    if not folder.is_dir():
        raise RuntimeError(f"Templates folder not found: {folder}")
    # Skip Word lock files (~$name.docx) and hidden files.
    files = sorted(p for p in folder.iterdir() if p.is_file() and not p.name.startswith(("~$", ".")))
    templates = []
    for path in files:
        text = read_template(path)
        if not text:
            log.warning("Template %s is empty or unsupported - skipped", path.name)
            continue
        first_line = next((l.lstrip("# ") for l in text.splitlines() if l.strip()), "")
        log.info("Template %s: %s", path.name, first_line[:80])
        templates.append((path.name, text))
    if not templates:
        raise RuntimeError(f"No templates in {folder}")
    if len(templates) != EXPECTED_TEMPLATES:
        log.warning("Expected %d templates (one per loan type), found %d in %s", EXPECTED_TEMPLATES, len(templates), folder)
    return templates


# ---------------------------------------------------------------- Source markdown

def is_step1_result(p: Path) -> bool:
    return p.suffix.lower() == ".md" and not p.name.endswith((ASSESSMENT_SUFFIX, REPLY_SUFFIX))


def find_source(arg: str | None) -> Path:
    if arg:
        p = Path(arg)
        if p.is_file():
            return p
        if p.is_dir():
            candidates = [c for c in p.glob("*.md") if is_step1_result(c)]
            if candidates:
                return max(candidates, key=lambda c: c.stat().st_mtime)
        raise RuntimeError(f"No step-1 markdown found at {p}")
    candidates = [c for c in Path(pipeline.OUTBOX_DIR).glob("*/*.md") if is_step1_result(c)]
    if not candidates:
        raise RuntimeError(f"No step-1 markdown found in {pipeline.OUTBOX_DIR}")
    return max(candidates, key=lambda c: c.stat().st_mtime)


# ---------------------------------------------------------------- Pipeline

def split_markers(answer: str) -> tuple[dict[str, str], str]:
    """Pull the leading 'REQUEST_TYPE: …' / 'TEMPLATE: …' lines off the answer."""
    lines = answer.strip().splitlines()
    markers, i = {}, 0
    while i < len(lines):
        if not lines[i].strip():
            i += 1
            continue
        # Tolerate markdown decoration the model may add around the markers (bold, code, heading).
        m = MARKER.match(re.sub(r"[*`#>]", "", lines[i]).strip())
        if not m:
            break
        markers[m.group(1).upper()] = m.group(2).strip()
        i += 1
    return markers, "\n".join(lines[i:]).strip()


def assess(source_md: Path, templates_dir: Path = TEMPLATES_DIR, reply: bool = True) -> Path:
    source_md = pipeline.long_path(source_md)
    log.info("Assessing %s", source_md.name)
    application = source_md.read_text(encoding="utf-8", errors="replace")
    templates = load_templates(templates_dir)

    prompt = build_assessment_prompt(application, templates)
    out_dir = source_md.parent
    # Save the prompt before the call, so a failed FOMO request still leaves it for inspection.
    (out_dir / "assessment_input.txt").write_text(prompt, encoding="utf-8")
    log.info("Sending assessment prompt to FOMO (%d chars)", len(prompt))
    answer = pipeline.call_fomo([{"text": prompt}], ASSESSMENT_MAX_TOKENS)

    markers, body = split_markers(answer)
    loan_type = markers.get("REQUEST_TYPE", "").lower()
    loan_type = loan_type if loan_type in LOAN_TYPES else None
    template = markers.get("TEMPLATE")
    template = template if template in {name for name, _ in templates} else None
    log.info("Request type: %s, template: %s",
             LOAN_TYPES[loan_type]["label"] if loan_type else "not recognised", template or "not recognised")

    out_file = out_dir / f"{source_md.stem}{ASSESSMENT_SUFFIX}"
    header = [
        f"# Credit assessment – {out_dir.name}",
        "",
        f"- **Request type:** {LOAN_TYPES[loan_type]['label'] if loan_type else 'not recognised'}",
        f"- **Template:** {f'`{template}`' if template else 'not recognised'}",
        f"- **Assessed:** {datetime.now():%Y-%m-%d %H:%M:%S}",
        f"- **Application:** `{source_md.name}`",
        f"- **Model:** `{pipeline.MODEL_ID}`",
        "",
        "---",
        "",
    ]
    out_file.write_text("\n".join(header) + body + "\n", encoding="utf-8")
    log.info("Saved %s", out_file)
    if reply:
        write_customer_reply(source_md)
    return out_file


# ---------------------------------------------------------------- Step 3: customer reply

def _is_section(line: str, name: str) -> bool:
    # Tolerate markdown decoration around the section labels (bold, heading, code).
    return re.sub(r"[*`#>\s]", "", line).upper() == f"{name}:"


def split_reply(answer: str) -> tuple[str | None, str, str]:
    """Split the step-3 answer into (letter type, notes for the advisor, letter)."""
    m = re.search(r"LETTER_TYPE\W*:\W*([a-z_]+)", answer, re.IGNORECASE)
    letter_type = m.group(1).lower() if m and m.group(1).lower() in LETTER_TYPES else None
    lines = answer.strip().splitlines()
    notes_at = next((i for i, l in enumerate(lines) if _is_section(l, "NOTES")), None)
    letter_at = next((i for i, l in enumerate(lines) if _is_section(l, "LETTER")), None)
    if letter_at is None:
        return letter_type, "", answer.strip()
    notes = "\n".join(lines[notes_at + 1:letter_at]).strip() if notes_at is not None and notes_at < letter_at else ""
    return letter_type, notes, "\n".join(lines[letter_at + 1:]).strip()


def write_customer_reply(source_md: Path) -> Path:
    source_md = pipeline.long_path(source_md)
    out_dir = source_md.parent
    assessment_file = out_dir / f"{source_md.stem}{ASSESSMENT_SUFFIX}"
    if not assessment_file.is_file():
        raise RuntimeError(f"No assessment to base the reply on: {assessment_file.name}")
    log.info("Writing customer reply for %s", source_md.name)
    application = source_md.read_text(encoding="utf-8", errors="replace")
    assessment = assessment_file.read_text(encoding="utf-8", errors="replace")

    prompt = build_customer_reply_prompt(application, assessment)
    (out_dir / "customer_reply_input.txt").write_text(prompt, encoding="utf-8")
    log.info("Sending customer reply prompt to FOMO (%d chars)", len(prompt))
    answer = pipeline.call_fomo([{"text": prompt}], REPLY_MAX_TOKENS)

    letter_type, notes, letter = split_reply(answer)
    log.info("Letter type: %s", letter_type or "not recognised")
    out_file = out_dir / f"{source_md.stem}{REPLY_SUFFIX}"
    header = [
        f"# Customer reply (DRAFT) – {out_dir.name}",
        "",
        f"- **Letter type:** {LETTER_TYPE_LABELS.get(letter_type, 'not recognised')}",
        f"- **Prepared:** {datetime.now():%Y-%m-%d %H:%M:%S}",
        f"- **Based on:** `{source_md.name}`, `{assessment_file.name}`",
        "",
        "> **Draft for advisor review.** Fill in the [bracketed] fields with data from the bank's systems "
        "(interest rate, APR, instalment, total cost) and check the content before sending.",
        "",
    ]
    if notes:
        header += ["## To fill in and check", "", notes, ""]
    header += ["---", "", ""]
    out_file.write_text("\n".join(header) + letter + "\n", encoding="utf-8")
    log.info("Saved %s", out_file)
    return out_file


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    parser = argparse.ArgumentParser(description="Step-1 markdown + 3 bank templates -> FOMO -> filled template with maximum loan estimate")
    parser.add_argument("source", nargs="?", help="step-1 .md file or outbox folder (default: newest in outbox)")
    parser.add_argument("--templates", default=str(TEMPLATES_DIR), metavar="FOLDER", help="folder with the template files")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--no-reply", action="store_true", help="skip the customer reply (step 3)")
    mode.add_argument("--reply-only", action="store_true", help="only (re)generate the customer reply from an existing assessment")
    args = parser.parse_args()
    source = find_source(args.source)
    if args.reply_only:
        write_customer_reply(source)
    else:
        assess(source, Path(args.templates), reply=not args.no_reply)


if __name__ == "__main__":
    main()
